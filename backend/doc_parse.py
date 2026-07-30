"""Parse supporting Excel models and context documents into plain text summaries for Claude."""
import os
import openpyxl
from pypdf import PdfReader


def summarize_xlsx(path: str, max_rows_per_sheet: int = 200) -> str:
    """Flatten an Excel workbook into a readable text table per sheet (values, not formulas)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    chunks = []
    for sheet in wb.worksheets:
        chunks.append(f"### Sheet: {sheet.title}")
        rows_written = 0
        for row in sheet.iter_rows(values_only=True):
            if all(cell is None for cell in row):
                continue
            cells = ["" if c is None else str(c) for c in row]
            chunks.append(" | ".join(cells))
            rows_written += 1
            if rows_written >= max_rows_per_sheet:
                chunks.append(f"... ({sheet.max_row - rows_written} more rows truncated)")
                break
    return "\n".join(chunks)


def extract_pdf_text(path: str, max_chars: int = 20000) -> str:
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    text = "\n".join(parts)
    return text[:max_chars]


def extract_context_doc(path: str) -> str:
    """Extract text from a context document: PDF or plain text."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_pdf_text(path)
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()[:20000]
