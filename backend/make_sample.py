"""Generate a synthetic sample deck, model, and speech clip for end-to-end testing.

Deliberately plants mismatches so we can verify the consistency/talk-track checks:
- Deck says revenue is $50M; the model actually shows $45M.
- Spoken talk track says $60M, matching neither the deck nor the model.
"""
import os
import wave

from pptx import Presentation
from pptx.util import Inches
import openpyxl
import win32com.client

OUT = os.path.join(os.path.dirname(__file__), "sample")
os.makedirs(OUT, exist_ok=True)


def make_deck():
    prs = Presentation()
    layout_title = prs.slide_layouts[0]
    layout_content = prs.slide_layouts[1]

    s1 = prs.slides.add_slide(layout_title)
    s1.shapes.title.text = "Acme Corp: FY24 Performance Review"
    s1.placeholders[1].text = "Prepared for the Executive Steering Committee"

    s2 = prs.slides.add_slide(layout_content)
    s2.shapes.title.text = "Revenue grew 20% year over year"
    s2.placeholders[1].text = "FY24 revenue reached $50M, up from $41.7M in FY23."

    s3 = prs.slides.add_slide(layout_content)
    s3.shapes.title.text = "We expanded into three new regions"
    s3.placeholders[1].text = "New market entry in APAC, EMEA, and LATAM drove incremental pipeline."

    path = os.path.join(OUT, "sample_deck.pptx")
    prs.save(path)
    return path


def make_model():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Fiscal Year", "Revenue ($M)", "YoY Growth"])
    ws.append(["FY23", 41.7, None])
    ws.append(["FY24", 45.0, "7.9%"])  # deliberately doesn't match deck's $50M / 20%
    path = os.path.join(OUT, "sample_model.xlsx")
    wb.save(path)
    return path


SPEECH_TEXT = (
    "Thanks everyone for joining. So, um, this year revenue actually hit sixty million dollars, "
    "which was a really strong result. Next slide. And, uh, we also expanded into three new regions "
    "which opened up a lot of new pipeline for next year."
)


def make_speech_wav():
    path = os.path.join(OUT, "sample_speech.wav")
    voice = win32com.client.Dispatch("SAPI.SpVoice")
    stream = win32com.client.Dispatch("SAPI.SpFileStream")
    stream.Open(path, 3)  # SSFMCreateForWrite
    voice.AudioOutputStream = stream
    voice.Speak(SPEECH_TEXT)
    stream.Close()
    return path


if __name__ == "__main__":
    print("Deck:", make_deck())
    print("Model:", make_model())
    print("Speech:", make_speech_wav())
