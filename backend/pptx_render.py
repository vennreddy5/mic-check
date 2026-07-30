"""Render a PPTX to per-slide PNGs via PowerPoint COM automation, and extract per-slide text."""
import os
import pythoncom
import win32com.client
from pptx import Presentation


def extract_slide_text(pptx_path: str) -> list[dict]:
    """Return [{index, title, body, notes}] for each slide, 0-indexed."""
    prs = Presentation(pptx_path)
    slides = []
    for i, slide in enumerate(prs.slides):
        title = ""
        body_parts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            is_title = shape == slide.shapes.title
            if is_title and not title:
                title = text
            else:
                body_parts.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({
            "index": i,
            "title": title,
            "body": "\n".join(body_parts),
            "notes": notes,
        })
    return slides


def render_slides_to_png(pptx_path: str, out_dir: str, width_px: int = 1280) -> list[str]:
    """Export each slide as a PNG using PowerPoint COM automation. Returns list of file paths, in order."""
    os.makedirs(out_dir, exist_ok=True)
    pptx_path = os.path.abspath(pptx_path)
    out_dir = os.path.abspath(out_dir)

    pythoncom.CoInitialize()
    powerpoint = None
    presentation = None
    try:
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        presentation = powerpoint.Presentations.Open(
            pptx_path, WithWindow=False
        )
        height_px = int(width_px * 9 / 16)
        presentation.Export(out_dir, "PNG", width_px, height_px)
    finally:
        if presentation is not None:
            presentation.Close()
        pythoncom.CoUninitialize()

    files = sorted(
        (f for f in os.listdir(out_dir) if f.lower().endswith(".png")),
        key=lambda f: int("".join(filter(str.isdigit, f)) or 0),
    )
    return [os.path.join(out_dir, f) for f in files]
